"""
附件文档解析（轻量版 v3.5）

PDF/Word/Excel/纯文本类附件 → 纯文本，注入当前对话上下文。
不做持久化存储；知识库化是后续独立功能。

限制：
- 单文件 ≤ 10MB（base64 解码后）
- 单文档注入上限 40k 字符（超长截断并标注）
"""

import base64
import io

from utils.logger import logger

MAX_DOC_BYTES = 10 * 1024 * 1024
MAX_DOC_TEXT_CHARS = 40_000

TEXT_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "csv",
    "json",
    "log",
    "py",
    "js",
    "ts",
    "tsx",
    "jsx",
    "html",
    "css",
    "sql",
    "yml",
    "yaml",
}


class DocumentParseError(Exception):
    """附件解析失败（消息可直接展示给用户）"""


def parse_document(*, filename: str, content_base64: str) -> str:
    """按扩展名解析文档为纯文本。失败抛 DocumentParseError。"""
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except Exception as exc:
        raise DocumentParseError(f"附件 {filename} 不是有效的文件内容") from exc

    if len(raw) > MAX_DOC_BYTES:
        raise DocumentParseError(f"附件 {filename} 超过 10MB 上限")
    if not raw:
        raise DocumentParseError(f"附件 {filename} 是空文件")

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "pdf":
            text = _parse_pdf(raw)
        elif ext == "docx":
            text = _parse_docx(raw)
        elif ext in ("xlsx", "xlsm"):
            text = _parse_xlsx(raw)
        elif ext in TEXT_EXTENSIONS:
            text = raw.decode("utf-8")
        else:
            raise DocumentParseError(
                f"暂不支持的文档格式：{filename}（支持 pdf / docx / xlsx / txt / md / csv 等）"
            )
    except DocumentParseError:
        raise
    except UnicodeDecodeError as exc:
        raise DocumentParseError(f"附件 {filename} 无法按文本解析（可能是二进制文件）") from exc
    except Exception as exc:
        logger.warning(f"[DOC_PARSE] {filename} 解析失败: {exc}")
        raise DocumentParseError(f"附件 {filename} 解析失败") from exc

    text = text.strip()
    if not text:
        raise DocumentParseError(f"附件 {filename} 未解析出文本内容（可能是扫描件/纯图片 PDF）")
    if len(text) > MAX_DOC_TEXT_CHARS:
        text = text[:MAX_DOC_TEXT_CHARS] + "\n\n…（内容过长，已截断）"
    return text


def _parse_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(f"【第 {i} 页】\n{t}" for i, t in enumerate(pages, 1) if t.strip())


def _parse_docx(raw: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(raw))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _parse_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    blocks = []
    for sheet in workbook.worksheets:
        lines = [f"=== Sheet: {sheet.title} ==="]
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)
